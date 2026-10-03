import io
import logging
from typing import Optional
import docx
import pypdf

logger = logging.getLogger("hiregenius.ai_ml_service.text_extraction")


class TextExtractionException(Exception):
    """Base exception for text extraction errors."""

    pass


class UnsupportedFileTypeError(TextExtractionException):
    """Raised when file extension or MIME type is not supported."""

    pass


class EmptyFileError(TextExtractionException):
    """Raised when uploaded file is empty (0 bytes)."""

    pass


class CorruptedFileError(TextExtractionException):
    """Raised when file contents are corrupted or cannot be parsed."""

    pass


class EncryptedFileError(TextExtractionException):
    """Raised when PDF file is password protected or encrypted."""

    pass


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract raw text from a PDF file in bytes using pypdf."""
    if not file_bytes:
        raise EmptyFileError("PDF file is empty (0 bytes).")

    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        if reader.is_encrypted:
            try:
                decrypted = reader.decrypt("")
                if decrypted == pypdf.PasswordResult.NOT_DECRYPTED:
                    raise EncryptedFileError(
                        "PDF is password-protected and cannot be extracted."
                    )
            except Exception:
                raise EncryptedFileError(
                    "PDF is password-protected and cannot be extracted."
                )

        text_pages = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_pages.append(page_text.strip())

        extracted_text = "\n\n".join(text_pages).strip()
        if not extracted_text:
            raise CorruptedFileError(
                "No extractable text found in PDF (file may be empty or contain non-OCR scanned images)."
            )

        return extracted_text

    except (EmptyFileError, EncryptedFileError):
        raise
    except pypdf.errors.PdfReadError as e:
        logger.warning(f"pypdf ReadError: {e}")
        raise CorruptedFileError(f"Corrupted or unreadable PDF document: {e}")
    except Exception as e:
        logger.warning(f"Unexpected PDF extraction error: {e}")
        raise CorruptedFileError(f"Failed to extract text from PDF: {e}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract raw text from a DOCX file in bytes using python-docx."""
    if not file_bytes:
        raise EmptyFileError("DOCX file is empty (0 bytes).")

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        extracted_sections = []

        # Extract text from paragraphs
        for p in doc.paragraphs:
            clean_text = p.text.strip()
            if clean_text:
                extracted_sections.append(clean_text)

        # Extract text from tables (common in CV/resume layouts)
        for table in doc.tables:
            for row in table.rows:
                cells_text = [
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                ]
                if cells_text:
                    extracted_sections.append(" | ".join(cells_text))

        extracted_text = "\n\n".join(extracted_sections).strip()
        if not extracted_text:
            raise CorruptedFileError(
                "No extractable text found in DOCX file."
            )

        return extracted_text

    except EmptyFileError:
        raise
    except Exception as e:
        logger.warning(f"DOCX extraction error: {e}")
        raise CorruptedFileError(f"Corrupted or unreadable DOCX document: {e}")


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Dispatcher that inspects file extension (.pdf vs .docx) and extracts text.

    Raises specific exceptions for unsupported extensions, empty files, or
    corrupted documents.
    """
    if not file_bytes:
        raise EmptyFileError("Uploaded file is empty (0 bytes).")

    clean_filename = (filename or "").strip().lower()

    if clean_filename.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif clean_filename.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    else:
        ext = clean_filename.rsplit(".", 1)[-1] if "." in clean_filename else "unknown"
        raise UnsupportedFileTypeError(
            f"Unsupported file type '.{ext}'. Only PDF (.pdf) and Word (.docx) documents are supported."
        )
